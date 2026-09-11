"""CSV and Excel export services using pandas and openpyxl."""

import io

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.crud import get_all_leads_for_export, get_processing_stats


def generate_leads_dataframe(db: Session) -> pd.DataFrame:
    """Extract all leads from database into a standardized pandas DataFrame."""
    leads = get_all_leads_for_export(db)
    records = []
    for lead in leads:
        created_str = (
            lead.created_at.strftime("%Y-%m-%d %H:%M:%S")
            if lead.created_at
            else ""
        )
        records.append({
            "ID": lead.id,
            "External ID": lead.external_id or "",
            "Name": lead.name,
            "Email": lead.email,
            "Company": lead.company or "",
            "Source": lead.source,
            "Notes": lead.notes or "",
            "Enrichment Status": lead.enrichment_status,
            "Enrichment Score": lead.enrichment_score if lead.enrichment_score is not None else "",
            "Enrichment Segment": lead.enrichment_segment or "",
            "Created At": created_str,
        })

    columns = [
        "ID",
        "External ID",
        "Name",
        "Email",
        "Company",
        "Source",
        "Notes",
        "Enrichment Status",
        "Enrichment Score",
        "Enrichment Segment",
        "Created At",
    ]

    if not records:
        return pd.DataFrame(columns=columns)
    return pd.DataFrame(records, columns=columns)


def export_leads_csv(db: Session) -> bytes:
    """Generate a clean UTF-8 encoded CSV string of all leads."""
    df = generate_leads_dataframe(db)
    output = io.StringIO()
    df.to_csv(output, index=False, encoding="utf-8")
    return output.getvalue().encode("utf-8")


def export_leads_excel(db: Session) -> bytes:
    """
    Generate a professionally formatted Excel workbook containing:
    1. 'Leads' worksheet with freeze panes, autofilter, bold headers, and column formatting.
    2. 'Summary' worksheet with core KPI metrics and source breakdown.
    """
    df_leads = generate_leads_dataframe(db)
    stats = get_processing_stats(db)

    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # Write Leads sheet
        df_leads.to_excel(writer, sheet_name="Leads", index=False)

        # Build Summary DataFrames
        kpi_data = [
            {"Metric": "Total Leads", "Value": stats["total_leads"]},
            {"Metric": "Total Webhooks", "Value": stats["total_webhooks_received"]},
            {"Metric": "Created", "Value": stats["created"]},
            {"Metric": "Duplicates", "Value": stats["duplicates"]},
            {"Metric": "Enrichment Success", "Value": stats["enrichment_success"]},
            {"Metric": "Enrichment Failed", "Value": stats["enrichment_failed"]},
        ]
        df_kpi = pd.DataFrame(kpi_data)
        df_kpi.to_excel(writer, sheet_name="Summary", startrow=0, index=False)

        source_breakdown = [
            {"Source": source, "Leads": count}
            for source, count in stats.get("leads_by_source", {}).items()
        ]
        if not source_breakdown:
            source_breakdown = [{"Source": "None", "Leads": 0}]
        df_sources = pd.DataFrame(source_breakdown)
        # Leave 2 blank rows after KPI table (6 rows + 1 header = row 7, startrow=9)
        start_row_sources = len(kpi_data) + 3
        df_sources.to_excel(
            writer,
            sheet_name="Summary",
            startrow=start_row_sources,
            index=False,
        )

        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="1F4E79", end_color="1F4E79", fill_type="solid"
        )
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        # Format Leads Sheet
        ws_leads = writer.sheets["Leads"]
        ws_leads.freeze_panes = "A2"
        if ws_leads.max_row >= 1 and ws_leads.max_column >= 1:
            ws_leads.auto_filter.ref = ws_leads.dimensions

        for col_idx in range(1, ws_leads.max_column + 1):
            cell = ws_leads.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Format rows and compute column widths for Leads
        for col_idx in range(1, ws_leads.max_column + 1):
            col_letter = get_column_letter(col_idx)
            col_header = ws_leads.cell(row=1, column=col_idx).value or ""
            max_len = len(str(col_header))

            for row_idx in range(2, ws_leads.max_row + 1):
                cell = ws_leads.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))

                # Specific column alignments and formats
                if col_header == "Enrichment Score" and isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0"
                    cell.alignment = Alignment(horizontal="right")
                elif col_header in ("ID", "Created At"):
                    cell.alignment = Alignment(horizontal="center")

            ws_leads.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Format Summary Sheet
        ws_summary = writer.sheets["Summary"]
        summary_header_fill = PatternFill(
            start_color="2F5597", end_color="2F5597", fill_type="solid"
        )

        # Style KPI Header
        for c in range(1, 3):
            cell = ws_summary.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = summary_header_fill
            cell.alignment = Alignment(horizontal="center")

        for r in range(2, len(kpi_data) + 2):
            for c in range(1, 3):
                ws_summary.cell(row=r, column=c).border = thin_border

        # Style Source breakdown Header
        source_header_row = start_row_sources + 1
        for c in range(1, 3):
            cell = ws_summary.cell(row=source_header_row, column=c)
            cell.font = header_font
            cell.fill = summary_header_fill
            cell.alignment = Alignment(horizontal="center")

        for r in range(source_header_row + 1, source_header_row + 1 + len(source_breakdown)):
            for c in range(1, 3):
                ws_summary.cell(row=r, column=c).border = thin_border

        ws_summary.column_dimensions["A"].width = 25
        ws_summary.column_dimensions["B"].width = 18

    return output.getvalue()
