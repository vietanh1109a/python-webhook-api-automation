"""Tests for CSV and Excel export endpoints."""

import io

import openpyxl
from fastapi.testclient import TestClient


def test_csv_export(client: TestClient) -> None:
    # Ingest a sample lead
    client.post(
        "/webhooks/leads",
        json={
            "external_id": "exp-1",
            "name": "Export Test",
            "email": "export@test.com",
            "company": "Export Co",
            "source": "website",
            "notes": "Testing CSV export",
        },
    )

    response = client.get("/api/export/csv")
    assert response.status_code == 200
    assert "text/csv" in response.headers["content-type"]
    assert "attachment; filename=" in response.headers["content-disposition"]

    content = response.text
    # Verify CSV headers and data
    assert "ID,External ID,Name,Email,Company,Source,Notes" in content
    assert "Export Test" in content
    assert "export@test.com" in content
    assert "Export Co" in content
    assert "website" in content


def test_excel_export_structure_and_sheets(client: TestClient) -> None:
    # Ingest leads
    client.post(
        "/webhooks/leads",
        json={
            "external_id": "xl-1",
            "name": "Excel User 1",
            "email": "user1@excel.com",
            "company": "Acme Tools",
            "source": "shopify",
            "notes": "Interested in bulk order",
        },
    )
    client.post(
        "/webhooks/leads",
        json={
            "name": "Excel User 2",
            "email": "user2@excel.com",
            "source": "website",
        },
    )

    response = client.get("/api/export/excel")
    assert response.status_code == 200
    assert "openxmlformats-officedocument" in response.headers["content-type"]
    assert "leads_export.xlsx" in response.headers["content-disposition"]

    # Load and inspect with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(response.content))

    # Verify sheet names
    assert "Leads" in wb.sheetnames
    assert "Summary" in wb.sheetnames

    # Verify Leads sheet contents
    ws_leads = wb["Leads"]
    header_values = [cell.value for cell in ws_leads[1]]
    expected_headers = [
        "ID",
        "External ID",
        "Name",
        "Email",
        "Company",
        "Source",
        "Notes",
        "Enrichment Status",
        "Enrichment Score",
        "Enrichment Segment",
        "Created At",
    ]
    assert header_values == expected_headers
    assert ws_leads.max_row >= 3  # Header + 2 rows
    assert ws_leads.freeze_panes == "A2"
    assert ws_leads.auto_filter.ref is not None

    # Verify Summary sheet contents
    ws_summary = wb["Summary"]
    summary_text = [str(cell.value) for row in ws_summary.iter_rows() for cell in row]
    assert "Total Leads" in summary_text
    assert "Total Webhooks" in summary_text
    assert "shopify" in summary_text
    assert "website" in summary_text
